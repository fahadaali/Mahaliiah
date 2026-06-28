import openpyxl, json, datetime, re
from hijri_converter import Hijri, Gregorian

TODAY = datetime.date(2026,6,28)
SRC = "/root/.claude/uploads/35c302de-df01-569d-b946-503c6f6a4b2c/fdbe1a86-_______________________.xlsx"
wb = openpyxl.load_workbook(SRC, data_only=True)

def to_date(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    return None

def parse_greg_str(s):
    """Parse a gregorian date that may appear as text dd/mm/yyyy or yyyy-mm-dd."""
    if s is None: return None
    if isinstance(s,(datetime.datetime,datetime.date)): return to_date(s)
    s=str(s).strip()
    m=re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', s)
    if m:
        d,mo,y=map(int,m.groups())
        try: return datetime.date(y,mo,d)
        except: return None
    m=re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m:
        y,mo,d=map(int,m.groups())
        try: return datetime.date(y,mo,d)
        except: return None
    return None

def parse_hijri(s):
    """Parse hijri dd/mm/yyyy text -> gregorian date."""
    if s is None: return None
    s=str(s).strip()
    m=re.match(r'^(\d{1,2})/(\d{1,2})/(\d{3,4})$', s)
    if m:
        d,mo,y=map(int,m.groups())
    else:
        m2=re.match(r'^(\d{3,4})-(\d{1,2})-(\d{1,2})$', s)
        if not m2: return None
        y,mo,d=map(int,m2.groups())
    if y<1300 or y>1500: return None
    try:
        return Hijri(y,mo,d).to_gregorian().to_date() if hasattr(Hijri(y,mo,d).to_gregorian(),'to_date') else datetime.date(*[Hijri(y,mo,d).to_gregorian().year, Hijri(y,mo,d).to_gregorian().month, Hijri(y,mo,d).to_gregorian().day])
    except Exception as e:
        return None

def status_of(d):
    if d is None: return "unknown"
    days=(d-TODAY).days
    if days<0: return "expired"
    if days<=90: return "soon"
    return "valid"

def days_left(d):
    if d is None: return None
    return (d-TODAY).days

def iso(d):
    return d.isoformat() if d else None

# ---------- LICENSES ----------
ws=wb['التراخيص ']
rows=list(ws.iter_rows(values_only=True))
licenses=[]
current_company=None
header_seen=False
for r in rows:
    cells=[ (str(c).strip() if c is not None else "") for c in r ]
    joined="".join(cells).strip()
    if not joined: continue
    c0,c1=cells[0],cells[1]
    # company line: c0 empty, c1 has a long company-ish name and not header
    if c0=="" and c1 and ("الشركة" in c1 or "مصنع" in c1):
        current_company=c1; continue
    if c1=="البيان": # header
        continue
    if "بيان السجلات" in joined and c0.startswith("بيان") or "بيان السجلات والتراخيص" in c0:
        continue
    # data row: has a name in c1
    name=c1
    if not name: continue
    if name in ("البيان",): continue
    hijri=cells[2]
    greg_raw=r[3] if len(r)>3 else None
    platform=cells[4]
    notes=cells[5] if len(cells)>5 else ""
    gd=to_date(greg_raw) or parse_greg_str(greg_raw)
    # if greg missing try hijri
    if gd is None and hijri:
        gd=parse_hijri(hijri)
    # text-y greg like "25/ من كل شهر"
    greg_text=""
    if gd is None and greg_raw:
        greg_text=str(greg_raw).strip()
    licenses.append({
        "company": current_company or "",
        "item": name,
        "hijri": hijri,
        "expiry": iso(gd),
        "expiry_text": greg_text,
        "platform": platform,
        "notes": notes,
        "status": status_of(gd) if gd else ("recurring" if "كل شهر" in greg_text else "unknown"),
        "days": days_left(gd)
    })

# ---------- IQAMAS ----------
ws=wb['بيان الاقامات ']
iqamas=[]
for i,r in enumerate(ws.iter_rows(values_only=True)):
    if i==0: continue
    if not r or not r[0]: continue
    name=str(r[0]).strip()
    company=str(r[1]).strip() if r[1] else ""
    num=str(r[2]).strip() if r[2] else ""
    exp=to_date(r[3]) or parse_greg_str(r[3])
    iqamas.append({"name":name,"company":company,"number":num,
        "expiry":iso(exp),"status":status_of(exp),"days":days_left(exp)})

# ---------- VEHICLES ----------
ws=wb['السيارات ']
vehicles=[]
for i,r in enumerate(ws.iter_rows(values_only=True)):
    if i==0: continue
    if not r or not r[0]: continue
    vtype=str(r[0]).strip()
    plate=str(r[1]).strip() if r[1] else ""
    company=str(r[2]).strip() if r[2] else ""
    form_h=str(r[3]).strip() if r[3] else ""
    insp_h=str(r[4]).strip() if r[4] else ""
    form_g=parse_hijri(form_h)
    insp_g=parse_hijri(insp_h)
    cond=str(r[6]).strip() if len(r)>6 and r[6] else ""
    vehicles.append({"type":vtype,"plate":plate,"company":company,
        "form_hijri":form_h,"form_expiry":iso(form_g),"form_status":status_of(form_g),"form_days":days_left(form_g),
        "insp_hijri":insp_h,"insp_expiry":iso(insp_g),"insp_status":status_of(insp_g) if insp_g else "unknown",
        "condition":cond})

# ---------- PLATFORMS ----------
ws=wb['سجل المنصات والانظمة ']
platforms=[]
for i,r in enumerate(ws.iter_rows(values_only=True)):
    if i==0: continue
    if not r or not r[0]: continue
    nm=str(r[0]).strip()
    if nm and nm!="أسم الجهة":
        platforms.append(nm)

data={"generated":TODAY.isoformat(),"licenses":licenses,"iqamas":iqamas,
      "vehicles":vehicles,"platforms":platforms}

with open("/home/user/Mahaliiah/data/data.json","w",encoding="utf-8") as f:
    json.dump(data,f,ensure_ascii=False,indent=1)

print("licenses",len(licenses),"iqamas",len(iqamas),"vehicles",len(vehicles),"platforms",len(platforms))
from collections import Counter
print("iqama status",Counter(x['status'] for x in iqamas))
print("vehicle form status",Counter(x['form_status'] for x in vehicles))
print("license status",Counter(x['status'] for x in licenses))
